"""Genera XLSX de despiece y cotización a partir del JSON de proyecto.

Uso:  python exportar_xlsx.py datos.json out_dir
Produce:
  - Despiece_<clave>.xlsx     (de datos["materiales"])
  - Cotizacion_<clave>.xlsx   (de datos["cotizacion"], si existe)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AZUL = "1A3A8C"
GRIS = "EFEFEF"


def _header(ws, titulo: str, subtitulo: str, ncols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, titulo)
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=AZUL)
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(2, 1, subtitulo).alignment = Alignment(horizontal="center")


def _anchos(ws, anchos: list[int]) -> None:
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fila_encabezado(ws, fila: int, cols: list[str]) -> None:
    for j, h in enumerate(cols, 1):
        c = ws.cell(fila, j, h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=GRIS)


def despiece(datos: dict, out_dir: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Despiece"
    cols = ["Pzas", "Código", "Descripción", "Color", "Obs."]
    _header(ws, f"DESPIECE — {datos.get('proyecto', '')}",
            f"{datos.get('clave', '')} · {datos.get('cliente', '')}", len(cols))
    _fila_encabezado(ws, 4, cols)
    r = 5
    for m in datos.get("materiales", []):
        ws.cell(r, 1, m.get("pzas"))
        ws.cell(r, 2, m.get("codigo"))
        ws.cell(r, 3, m.get("descripcion"))
        ws.cell(r, 4, m.get("color"))
        ws.cell(r, 5, m.get("obs", ""))
        r += 1
    _anchos(ws, [8, 18, 50, 16, 26])
    p = out_dir / f"Despiece_{datos.get('clave', 'PROYECTO')}.xlsx"
    wb.save(p)
    return p


def cotizacion(datos: dict, out_dir: Path) -> Path | None:
    items = datos.get("cotizacion") or []
    if not items:
        # Derivar de los materiales que traigan precio unitario.
        items = [
            {
                "codigo": m.get("codigo"),
                "descripcion": m.get("descripcion"),
                "cantidad": m.get("pzas"),
                "precio_unitario": m.get("precio"),
            }
            for m in datos.get("materiales", [])
            if m.get("precio") is not None
        ]
    if not items:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Cotización"
    cols = ["Código", "Descripción", "Cant.", "P. Unit.", "Importe"]
    _header(ws, f"COTIZACIÓN — {datos.get('proyecto', '')}",
            f"{datos.get('clave', '')} · {datos.get('cliente', '')} · MXN sin IVA (mayoreo)",
            len(cols))
    _fila_encabezado(ws, 4, cols)
    r = 5
    subtotal = 0.0
    for it in items:
        imp = it.get("importe")
        if imp is None:
            try:
                imp = float(it.get("cantidad", 0)) * float(it.get("precio_unitario", 0))
            except (TypeError, ValueError):
                imp = 0
        subtotal += float(imp or 0)
        ws.cell(r, 1, it.get("codigo"))
        ws.cell(r, 2, it.get("descripcion"))
        ws.cell(r, 3, it.get("cantidad"))
        cu = ws.cell(r, 4, it.get("precio_unitario"))
        cu.number_format = "#,##0.00"
        ci = ws.cell(r, 5, imp)
        ci.number_format = "#,##0.00"
        r += 1
    ws.cell(r, 4, "Subtotal").font = Font(bold=True)
    cs = ws.cell(r, 5, round(subtotal, 2))
    cs.font = Font(bold=True)
    cs.number_format = "#,##0.00"
    ws.cell(r + 2, 2, "IVA, flete e instalación se cotizan por separado.").font = Font(italic=True)
    _anchos(ws, [18, 54, 10, 14, 16])
    p = out_dir / f"Cotizacion_{datos.get('clave', 'PROYECTO')}.xlsx"
    wb.save(p)
    return p


def generar(datos: dict, out_dir) -> list[Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    salidas = [despiece(datos, out_dir)]
    cot = cotizacion(datos, out_dir)
    if cot:
        salidas.append(cot)
    return salidas


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Uso: python exportar_xlsx.py datos.json out_dir")
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        datos = json.load(f)
    for p in generar(datos, sys.argv[2]):
        print("XLSX:", p)
